"""Exportación Excel de operaciones del backtest RSI + media RSI."""
from datetime import datetime, timezone
from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from risk_management import RiskConfig, calculate_position, scenario_pnl_usdt
from indicators import rsi_indicator

DARK=PatternFill('solid', fgColor='1F2937'); WHITE=Font(bold=True, color='FFFFFF'); BOLD=Font(bold=True); TITLE=Font(bold=True, size=14)

def build_rsi_workbook(symbol, interval, horizon, klines, trades, stats, risk_config: RiskConfig,
                       rsi_period=14, ma_type='SMA', ma_period=14, sl_lookback=10, stop_mode='lookback', pivot_left=2, pivot_right=2, pivot_buffer_pct=0.1, ignore_first_ma_cross=False, use_extreme_filter=False):
    wb=Workbook(); ws=wb.active; ws.title='Resumen'
    ws['A1']=f'RSI + Media RSI — {symbol} {interval}'; ws['A1'].font=TITLE; ws.merge_cells('A1:D1')
    first=datetime.fromtimestamp(int(klines[0]['time'])/1000,timezone.utc); last=datetime.fromtimestamp(int(klines[-1]['time'])/1000,timezone.utc)
    meta=[('Símbolo',symbol),('Temporalidad',interval),('Período UTC',f'{first:%Y-%m-%d %H:%M} → {last:%Y-%m-%d %H:%M}'),
      ('Velas cargadas',len(klines)),('Horizonte',horizon if horizon is not None else 'Sin límite'),('RSI',f'Wilder/RMA, período {rsi_period}'),
      ('Media RSI',f'{ma_type}, período {ma_period}'),('Regímenes','RSI >= 70: solo SHORT; RSI <= 30: solo LONG; se mantiene hasta el extremo opuesto'),
      ('Entradas','Primera por giro del RSI; adicionales por cruce RSI/media con LONG RSI < 55 y SHORT RSI > 45'),
      ('Omitir primer cruce de media', 'Sí' if ignore_first_ma_cross else 'No; contar todos los cruces válidos'),
      ('Reinicio de régimen','Al volver a entrar en el mismo extremo tras salir de la zona, se reinician giro inicial y contador de cruces; las operaciones abiertas siguen activas'),
      ('Entrada','Apertura de la vela siguiente'),('Stop', ('Último pivote local confirmado con colchón '+str(pivot_buffer_pct)+'% (izquierda='+str(pivot_left)+', derecha='+str(pivot_right)+'); si no existe, extremo de ventana' if stop_mode == 'pivot' else 'Mínimo/máximo local de las últimas '+str(sl_lookback)+' velas, incluida señal')),
      ('Modo de stop',stop_mode),
      ('Colchón del pivote %',pivot_buffer_pct),
      ('Objetivo','2R fijo'),('Posiciones simultáneas','Permitidas varias; cada operación se simula de forma independiente'),('Riesgo',risk_config.label),('Costes','No incluye comisiones, funding ni deslizamiento')]
    row=3
    for k,v in meta: ws.cell(row,1,k).font=BOLD; ws.cell(row,2,v); row+=1
    row+=1
    for c,h in enumerate(['Métrica','Total','LONG','SHORT'],1): ws.cell(row,c,h).fill=DARK; ws.cell(row,c).font=WHITE
    metrics=[('Operaciones','total'),('Resueltas','resueltas'),('TP','tp'),('SL','sl'),('Horizonte','horizonte'),('Pendientes','pendientes'),('Acierto %','win_rate'),('R netas','r_total'),('R media','r_medio'),('Profit factor','profit_factor'),('Drawdown máximo R','max_drawdown_r'),('Cambio % acumulado','pnl_pct_total')]
    for ri,(label,key) in enumerate(metrics,row+1):
        ws.cell(ri,1,label)
        for col,g in enumerate(['total','long','short'],2):
            val=stats[g][key]; ws.cell(ri,col,'∞' if isinstance(val,float) and val==float('inf') else val)
    for c,w in enumerate([34,24,24,24],1): ws.column_dimensions[get_column_letter(c)].width=w
    op=wb.create_sheet('Operaciones')
    headers=['Hora señal UTC','Hora entrada UTC','Tipo de señal','Lado','Entrada','Stop','Objetivo 2R','RSI señal','Media RSI señal','Riesgo %','Objetivo R','Resultado','Hora salida UTC','Precio salida','Resultado R','Cambio %','Riesgo USDT','Cantidad estimada','Nominal USDT','Margen aislado USDT','Apalancamiento','PnL USDT estimado']
    for c,h in enumerate(headers,1): op.cell(1,c,h).fill=DARK; op.cell(1,c).font=WHITE
    for r,t in enumerate(trades,2):
        pos=calculate_position(t.entry,t.stop_loss,risk_config); pnl=scenario_pnl_usdt(t.r_multiple,pos) if t.r_multiple is not None else None
        vals=[datetime.fromtimestamp(t.time/1000,timezone.utc).replace(tzinfo=None),datetime.fromtimestamp(t.entry_time/1000,timezone.utc).replace(tzinfo=None),getattr(t,'signal_type','MA_CROSS'),t.side,t.entry,t.stop_loss,t.target,t.rsi,t.rsi_ma,t.risk_pct,t.target_r,t.outcome,datetime.fromtimestamp(t.exit_time/1000,timezone.utc).replace(tzinfo=None) if t.exit_time else None,t.exit_price,t.r_multiple,t.pnl_pct,pos.risk_usdt,pos.quantity,pos.notional_usdt,pos.isolated_margin_usdt,pos.leverage,pnl]
        for c,v in enumerate(vals,1): op.cell(r,c,v)
    op.freeze_panes='A2'; op.auto_filter.ref=op.dimensions
    for c in range(1,len(headers)+1): op.column_dimensions[get_column_letter(c)].width=18
    for r in range(2,op.max_row+1):
        for c in (1,2,12): op.cell(r,c).number_format='yyyy-mm-dd hh:mm'
    candles=wb.create_sheet('Velas')
    candle_closes=[float(k['close']) for k in klines]
    candle_volumes=[float(k.get('volume',0) or 0) for k in klines]
    rsi_data=rsi_indicator(candle_closes, rsi_period=rsi_period, ma_type=ma_type, ma_period=ma_period, volume=candle_volumes)
    headers=['Hora UTC','Open','High','Low','Close','Volumen','RSI','Media RSI','Banda RSI superior','Banda RSI inferior']
    for c,h in enumerate(headers,1): candles.cell(1,c,h).fill=DARK; candles.cell(1,c).font=WHITE
    for i,k in enumerate(klines):
        r=i+2
        vals=[datetime.fromtimestamp(int(k['time'])/1000,timezone.utc).replace(tzinfo=None),float(k['open']),float(k['high']),float(k['low']),float(k['close']),candle_volumes[i],rsi_data['rsi'][i],rsi_data['ma'][i],rsi_data['bb_upper'][i],rsi_data['bb_lower'][i]]
        for c,v in enumerate(vals,1): candles.cell(r,c,v)
        candles.cell(r,1).number_format='yyyy-mm-dd hh:mm'
        for c in range(7,11):
            candles.cell(r,c).number_format='0.00'
    candles.freeze_panes='A2'; candles.auto_filter.ref=candles.dimensions
    for c,w in enumerate([20,16,16,16,16,18,12,14,20,20],1): candles.column_dimensions[get_column_letter(c)].width=w
    out=BytesIO(); wb.save(out); out.seek(0); return out
