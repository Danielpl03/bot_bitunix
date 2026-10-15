"""Exportación Excel del backtest de Keltner Channel."""
from __future__ import annotations

from datetime import datetime, timezone
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from backtest import OUTCOME_NA, OUTCOME_PENDING, OUTCOME_SL, OUTCOME_TP, ratio_label
from backtest_kelner import KELTNER_TP_RATIOS, KeltnerBacktestSignal
from risk_management import RiskConfig, calculate_position, scenario_pnl_usdt

FONT_NAME = "Arial"
HEADER_FILL = PatternFill("solid", fgColor="1F2937")
HEADER_FONT = Font(name=FONT_NAME, bold=True, color="FFFFFF")
TITLE_FONT = Font(name=FONT_NAME, bold=True, size=14)
LABEL_FONT = Font(name=FONT_NAME, bold=True)
LONG_FILL = PatternFill("solid", fgColor="DCFCE7")
SHORT_FILL = PatternFill("solid", fgColor="FEE2E2")
SECTION_FILL = PatternFill("solid", fgColor="E5E7EB")


def _dt(ms: int) -> datetime:
    return datetime.fromtimestamp(int(ms) / 1000, tz=timezone.utc).replace(tzinfo=None)


def _header(ws, row: int, n: int) -> None:
    for col in range(1, n + 1):
        c = ws.cell(row, col)
        c.font, c.fill = HEADER_FONT, HEADER_FILL
        c.alignment = Alignment(horizontal="center", vertical="center")


def _autosize(ws, widths: list[int]) -> None:
    for i, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = width


def _summary(ws, symbol, interval, horizon, klines, stats):
    ws.title = "Resumen"
    ws["A1"] = f"Backtest Keltner Channel — {symbol} — {interval}"
    ws["A1"].font = TITLE_FONT
    ws.merge_cells("A1:D1")
    meta = [
        ("Estrategia", "Cruce de bandas Keltner"),
        ("Símbolo", symbol), ("Intervalo", interval),
        ("Límite de velas por operación", horizon if horizon is not None else "Sin límite"),
        ("Velas analizadas", len(klines)), ("Desde (UTC)", _dt(klines[0]["time"]).strftime("%Y-%m-%d %H:%M")),
        ("Hasta (UTC)", _dt(klines[-1]["time"]).strftime("%Y-%m-%d %H:%M")),
        ("Generado (UTC)", datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M")),
    ]
    row = 3
    for label, value in meta:
        ws.cell(row, 1, label).font = LABEL_FONT
        ws.cell(row, 2, value)
        row += 1
    row += 1
    for col, value in enumerate(["Métrica", "Total", "LONG", "SHORT"], 1):
        ws.cell(row, col, value)
    _header(ws, row, 4)
    row += 1
    groups = (("total", 2), ("long", 3), ("short", 4))
    for label, key in [("Señales totales", "total"), ("Riesgo medio (%)", "riesgo_medio_pct")]:
        ws.cell(row, 1, label).font = LABEL_FONT
        for group, col in groups:
            value = stats[group][key]
            ws.cell(row, col, round(value, 4) if value is not None else "—")
        row += 1
    for ratio in KELTNER_TP_RATIOS:
        for col in range(1, 5):
            ws.cell(row, col).fill = SECTION_FILL
        ws.cell(row, 1, f"Escenario objetivo {ratio_label(ratio)}").font = LABEL_FONT
        row += 1
        metrics = [
            ("Operaciones resueltas", "resueltas"), ("Llegan al objetivo", "tp"),
            ("Stop loss", "sl"), ("Cerradas al horizonte", "horizonte"),
            ("Pendientes", "pendientes"), ("Win rate (%)", "win_rate"),
            ("R total", "r_total"), ("R medio", "r_medio"),
            ("Cambio total (%)", "cambio_total_pct"), ("Cambio medio (%)", "cambio_medio_pct"),
        ]
        for label, key in metrics:
            ws.cell(row, 1, label).font = LABEL_FONT
            for group, col in groups:
                value = stats[group]["escenarios"][ratio][key]
                ws.cell(row, col, round(value, 4) if value is not None else "—")
            row += 1
    for col in range(1, 5):
        ws.cell(row, col).fill = SECTION_FILL
    ws.cell(row, 1, "Ratio máximo alcanzado (R)").font = LABEL_FONT
    row += 1
    for label, key in [("Medido en operaciones", "medidas"), ("Operaciones abiertas (medición provisional)", "abiertas"),
                       ("Media", "medio"), ("Mediana", "mediana"), ("Máximo", "maximo")]:
        ws.cell(row, 1, label).font = LABEL_FONT
        for group, col in groups:
            value = stats[group]["max_ratio"][key]
            ws.cell(row, col, round(value, 4) if value is not None else "—")
        row += 1
    _autosize(ws, [52, 22, 22, 22])
    ws.freeze_panes = "A3"


def _signals(ws, results):
    ws.title = "Señales"
    columns = [
        ("Fecha UTC", 18, "date"), ("Lado", 9, ""), ("Apertura", 14, "price"),
        ("Entrada/cierre", 16, "price"), ("Banda inferior", 15, "price"),
        ("Banda media", 15, "price"), ("Banda superior", 15, "price"),
        ("Bandas cruzadas", 20, ""), ("Stop loss", 14, "price"), ("Riesgo %", 11, "pct"),
        ("Ratio máximo (R)", 15, "r"), ("Precio máximo favorable", 20, "price"),
        ("Velas hasta máximo", 17, ""), ("Fin recorrido", 16, ""), ("R al horizonte", 14, "r"),
    ]
    for ratio in KELTNER_TP_RATIOS:
        lab = ratio_label(ratio)
        columns += [(f"Objetivo {lab}", 15, "price"), (f"Resultado {lab}", 15, ""),
                    (f"Salida {lab}", 15, "price"), (f"Cambio % {lab}", 15, "pct_signed"),
                    (f"R {lab}", 10, "r")]
    for col, (name, _, _) in enumerate(columns, 1):
        ws.cell(1, col, name)
    _header(ws, 1, len(columns))
    for i, sig in enumerate(results, 2):
        vals = [_dt(sig.time), sig.side, sig.open, sig.close, sig.lower, sig.middle, sig.upper,
                ", ".join(sig.crossed_bands), sig.stop_loss, sig.risk_pct, sig.max_ratio,
                sig.max_ratio_price, sig.candles_to_max, sig.trade_end, sig.horizon_r]
        for ratio in KELTNER_TP_RATIOS:
            sc = sig.scenarios[ratio]
            vals += [sc.target, sc.outcome, sc.exit_price, sc.pct_change, sc.r_multiple]
        for col, value in enumerate(vals, 1):
            cell = ws.cell(i, col, value)
            fmt = columns[col-1][2]
            if fmt == "date": cell.number_format = "yyyy-mm-dd hh:mm"
            elif fmt == "price" and isinstance(value, (int, float)): cell.number_format = "0.000000"
            elif fmt == "pct" and value is not None: cell.number_format = '0.00"%"'
            elif fmt == "pct_signed" and value is not None: cell.number_format = '+0.00"%";-0.00"%"'
            elif fmt == "r" and value is not None: cell.number_format = '+0.00"R";-0.00"R"'
        ws.cell(i, 2).fill = LONG_FILL if sig.side == "LONG" else SHORT_FILL
        for j, ratio in enumerate(KELTNER_TP_RATIOS):
            outcome = sig.scenarios[ratio].outcome
            if outcome == OUTCOME_TP: fill = LONG_FILL
            elif outcome == OUTCOME_SL: fill = SHORT_FILL
            else: fill = None
            if fill:
                base = 16 + j * 5
                for col in (base+1, base+2): ws.cell(i, col).fill = fill
    _autosize(ws, [w for _, w, _ in columns])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{len(results)+1}"


def _candles(ws, klines):
    ws.title = "Velas"
    for col, name in enumerate(["Fecha UTC", "Open", "High", "Low", "Close"], 1):
        ws.cell(1, col, name)
    _header(ws, 1, 5)
    for i, k in enumerate(klines, 2):
        for col, value in enumerate([_dt(k["time"]), float(k["open"]), float(k["high"]),
                                     float(k["low"]), float(k["close"])], 1):
            c = ws.cell(i, col, value)
            c.number_format = "yyyy-mm-dd hh:mm" if col == 1 else "0.000000"
    _autosize(ws, [18, 14, 14, 14, 14])
    ws.freeze_panes = "A2"


def _curve(ws, stats):
    ws.title = "Curva ratios"
    headers = ["Grupo", "Ratio TP (R)", "Operaciones", "Alcanzan", "% alcanzan", "R total", "R medio"]
    for col, name in enumerate(headers, 1): ws.cell(1, col, name)
    _header(ws, 1, len(headers))
    row = 2
    for group, label in [("total", "Total"), ("long", "LONG"), ("short", "SHORT")]:
        for item in stats[group]["curva"]:
            vals = [label, item["ratio"], item["operaciones"], item["alcanzan"],
                    item["pct_alcanzan"], item["r_total"], item["r_medio"]]
            for col, value in enumerate(vals, 1):
                c = ws.cell(row, col, value)
                if col in (5,): c.number_format = '0.00"%"'
                if col in (6, 7): c.number_format = '+0.00"R";-0.00"R"'
            row += 1
    _autosize(ws, [14, 15, 15, 13, 15, 14, 14])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:G{row-1}"



def _risk_sheet(ws, results, config: RiskConfig):
    ws.title = "Gestión riesgo"
    ws["A1"] = "Dimensionamiento de posiciones — estimación"
    ws["A1"].font = TITLE_FONT
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=14)
    ws["A2"] = "Modalidad"; ws["B2"] = config.label
    ws["A3"] = "Margen"; ws["B3"] = "Aislado"
    ws["A4"] = "Buffer liquidación"; ws["B4"] = config.safety_buffer_percent / 100; ws["B4"].number_format = '0.00%'
    ws["A5"] = "Aviso"; ws["B5"] = "Estimación; confirmar tier, mantenimiento y apalancamiento máximo real de Bitunix."
    headers = ["Fecha UTC", "Lado", "Entrada", "Stop", "Distancia stop %", "Riesgo USDT", "Cantidad",
               "Nominal USDT", "Apalancamiento x", "Margen aislado USDT", "Distancia liquidación estimada %", "Cumple buffer", "P/L USDT", "Nota"]
    start = 7
    for col, name in enumerate(headers, 1): ws.cell(start, col, name)
    _header(ws, start, len(headers))
    row = start + 1
    for sig in results:
        pos = calculate_position(sig.close, sig.stop_loss, config)
        for ratio in KELTNER_TP_RATIOS:
            sc = sig.scenarios[ratio]
            pnl = scenario_pnl_usdt(sc.r_multiple, pos)
            vals = [_dt(sig.time), f"{sig.side} · TP {ratio:g}:1", sig.close, sig.stop_loss, pos.stop_distance_percent,
                    pos.risk_usdt, pos.quantity, pos.notional_usdt, pos.leverage, pos.isolated_margin_usdt,
                    pos.estimated_liquidation_distance_percent, "Sí" if pos.feasible else "NO — descartar", pnl, pos.note]
            for col, value in enumerate(vals, 1):
                c = ws.cell(row, col, value)
                if col == 1: c.number_format = "yyyy-mm-dd hh:mm"
                elif col in (3, 4, 7, 8, 10): c.number_format = "0.000000"
                elif col in (5, 11): c.number_format = '0.00"%"'
                elif col in (6, 13): c.number_format = '0.00" USDT"'
            row += 1
    _autosize(ws, [18, 17, 14, 14, 16, 14, 16, 17, 16, 19, 24, 18, 15, 70])
    ws.freeze_panes = f"A{start+1}"
    ws.auto_filter.ref = f"A{start}:N{row-1}"


def build_keltner_workbook(symbol: str, interval: str, horizon: int | None,
                           klines: list[dict], results: list[KeltnerBacktestSignal],
                           stats: dict, risk_config: RiskConfig | None = None, signal_filter: dict | None = None) -> BytesIO:
    """Construye un Excel con resumen, señales, velas y curva de ratios."""
    wb = Workbook()
    _summary(wb.active, symbol, interval, horizon, klines, stats)
    _signals(wb.create_sheet("Señales"), results)
    _candles(wb.create_sheet("Velas"), klines)
    _curve(wb.create_sheet("Curva ratios"), stats)
    _risk_sheet(wb.create_sheet("Gestión riesgo"), results, risk_config or RiskConfig())
    filter_ws = wb.create_sheet("Filtro señales")
    enabled = bool(signal_filter and signal_filter.get("window") is not None)
    filter_ws.append(["Parámetro", "Valor"])
    filter_ws.append(["Filtro de agrupación", "Activado" if enabled else "Desactivado"])
    if enabled:
        filter_ws.append(["Ventana de velas anteriores", signal_filter["window"]])
        filter_ws.append(["Límite de señales previas", signal_filter["limit"]])
        filter_ws.append(["Señales rechazadas por agrupación dentro del período", signal_filter.get("rejected", 0)])
        filter_ws.append(["Regla de agrupación", "Se rechaza la señal actual si en las X velas previas ya hubo al menos el límite; las rechazadas también cuentan para señales futuras."])
    band_ratio = signal_filter.get("band_ratio") if signal_filter else None
    filter_ws.append(["Filtro de banda contraria", "Activado" if band_ratio is not None else "Desactivado"])
    if band_ratio is not None:
        filter_ws.append(["Ratio mínimo de espacio hasta banda contraria", band_ratio])
        filter_ws.append(["Señales rechazadas por banda dentro del período", signal_filter.get("band_rejected", 0)])
        filter_ws.append(["Regla de banda", "LONG: Upper - entrada > ratio × riesgo. SHORT: entrada - Lower > ratio × riesgo. La desigualdad es estricta."])
    only_outer_bands = bool(signal_filter and signal_filter.get("only_outer_bands"))
    filter_ws.append(["Entradas solo en bandas límite", "Activado" if only_outer_bands else "Desactivado"])
    filter_ws.append(["Regla de entrada", "LONG: cruce al alza de la banda inferior; SHORT: cruce a la baja de la banda superior." if only_outer_bands else "Keltner: se permite cruce al alza de banda inferior o media para LONG y cruce a la baja de banda superior o media para SHORT."])
    width_pct = signal_filter.get("channel_width_pct") if signal_filter else None
    filter_ws.append(["Filtro de anchura del canal", "Activado" if width_pct is not None else "Desactivado"])
    if width_pct is not None:
        filter_ws.append(["Anchura mínima del canal (% del precio de entrada)", width_pct])
        filter_ws.append(["Señales rechazadas por anchura dentro del período", signal_filter.get("width_rejected", 0)])
        filter_ws.append(["Regla de anchura", "Se acepta si (Upper - Lower) / precio de entrada × 100 es mayor o igual al mínimo configurado."])
    filter_ws.column_dimensions["A"].width = 45
    filter_ws.column_dimensions["B"].width = 100
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if cell.font is None or cell.font.name != FONT_NAME:
                    old = cell.font
                    cell.font = Font(name=FONT_NAME, bold=old.bold if old else False,
                                     color=old.color if old else None, size=old.size if old else 11)
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
