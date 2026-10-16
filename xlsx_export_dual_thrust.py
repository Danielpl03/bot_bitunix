"""Exportación Excel para Dual Thrust."""
from datetime import datetime, timezone
from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

DARK = PatternFill("solid", fgColor="1F2937")
WHITE = Font(bold=True, color="FFFFFF")
BOLD = Font(bold=True)
TITLE = Font(bold=True, size=14)


def build_dual_thrust_workbook(symbol, interval, horizon, klines, trades, stats, lookback, k1, k2):
    wb = Workbook(); ws = wb.active; ws.title = "Resumen"
    ws["A1"] = f"Dual Thrust Breakout — {symbol} {interval}"; ws["A1"].font = TITLE; ws.merge_cells("A1:D1")
    first = datetime.fromtimestamp(int(klines[0]["time"])/1000, timezone.utc)
    last = datetime.fromtimestamp(int(klines[-1]["time"])/1000, timezone.utc)
    meta = [
        ("Símbolo", symbol), ("Temporalidad", interval),
        ("Período UTC", f"{first:%Y-%m-%d %H:%M} → {last:%Y-%m-%d %H:%M}"),
        ("Velas cargadas", len(klines)), ("Horizonte externo", horizon if horizon else "Sin límite"),
        ("Ventana histórica N", lookback), ("Coeficiente alcista K1", k1), ("Coeficiente bajista K2", k2),
        ("Fórmula Range", "max(HH - LC, HC - LL) usando las N velas anteriores"),
        ("Niveles", "BuyTrack = Open + K1×Range; SellTrack = Open − K2×Range"),
        ("Ejecución", "Ruptura intrabar; reversión inmediata ante ruptura opuesta"),
        ("Barras ambiguas", "Si High alcanza BuyTrack y Low alcanza SellTrack en la misma vela, se ignora la señal"),
        ("Trailing stop", "Desactivado en la versión base; opcional en la investigación"),
        ("Costes", "No incluye comisiones, financiación ni deslizamiento"),
        ("Riesgo", "No se calcula tamaño por riesgo porque la versión base no define stop-loss"),
    ]
    row = 3
    for key, value in meta:
        ws.cell(row, 1, key).font = BOLD; ws.cell(row, 2, value); row += 1
    row += 1
    for c, label in enumerate(["Métrica", "Total", "LONG", "SHORT"], 1):
        ws.cell(row, c, label).fill = DARK; ws.cell(row, c).font = WHITE
    metrics = [("Operaciones", "total"), ("Resueltas", "resueltas"), ("Reversiones", "reversiones"),
               ("Cerradas por horizonte", "horizonte"), ("Pendientes", "pendientes"), ("Acierto %", "win_rate"),
               ("PnL % acumulado (suma)", "pnl_pct_total"), ("PnL % medio", "pnl_pct_medio"),
               ("Profit factor (%)", "profit_factor"), ("Drawdown máximo % (sobre suma de retornos)", "max_drawdown_pct")]
    for ri, (label, key) in enumerate(metrics, row+1):
        ws.cell(ri, 1, label)
        for col, group in enumerate(["total", "long", "short"], 2):
            val = stats[group][key]
            ws.cell(ri, col, "∞" if isinstance(val, float) and val == float("inf") else val)
    for col, width in enumerate([42, 25, 25, 25], 1): ws.column_dimensions[get_column_letter(col)].width = width
    op = wb.create_sheet("Operaciones")
    headers = ["Hora entrada UTC", "Hora salida UTC", "Lado", "Entrada", "Salida", "Motivo salida", "Range", "BuyTrack entrada", "SellTrack entrada", "Velas mantenida", "PnL %"]
    for c, h in enumerate(headers, 1): op.cell(1, c, h).fill = DARK; op.cell(1, c).font = WHITE
    for r, t in enumerate(trades, 2):
        vals = [datetime.fromtimestamp(t.entry_time/1000, timezone.utc).replace(tzinfo=None),
                datetime.fromtimestamp(t.exit_time/1000, timezone.utc).replace(tzinfo=None) if t.exit_time else None,
                t.side, t.entry, t.exit_price, t.exit_reason, t.range_value, t.buy_track_entry, t.sell_track_entry, t.bars_held, t.pnl_pct]
        for c, value in enumerate(vals, 1): op.cell(r, c, value)
    op.freeze_panes = "A2"; op.auto_filter.ref = op.dimensions
    for c in range(1, len(headers)+1): op.column_dimensions[get_column_letter(c)].width = 20
    for r in range(2, op.max_row+1):
        op.cell(r, 1).number_format = "yyyy-mm-dd hh:mm"; op.cell(r, 2).number_format = "yyyy-mm-dd hh:mm"
    out = BytesIO(); wb.save(out); out.seek(0); return out
