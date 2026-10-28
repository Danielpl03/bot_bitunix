from xlsx_export_ttr import build_ttr_workbook
from openpyxl import load_workbook


def test_ttr_export_contains_partial_pnl_breakdown():
    trade = {
        'side': 'sell', 'entry_time': 1000, 'exit_time': 2000,
        'entry_price': 100.0, 'avg_entry_price': 100.0, 'exit_price': 100.0,
        'exit_reason': 'SL_PSAR', 'pnl_pct': 0.4, 'remaining_pct': 0.0,
        'bars_held': 1,
        'partial_exits': [
            {'index': 1, 'time': 1500, 'price': 99.0, 'percent': 20.0, 'reason': 'TP', 'pnl_pct_position': 0.4, 'remaining_pct': 80.0},
            {'index': 2, 'time': 2000, 'price': 100.0, 'percent': 80.0, 'reason': 'SL_PSAR', 'pnl_pct_position': 0.0, 'remaining_pct': 0.0},
        ],
    }
    result = {'summary': {}, 'events': [], 'parameters': {}, 'note': 'test'}
    data = build_ttr_workbook('BTCUSDT', '15m', [{'time': 1000}, {'time': 2000}], result, [trade])
    wb = load_workbook(data, data_only=True)
    assert {'Resumen', 'Operaciones', 'Cierres parciales', 'Señales', 'Curva de capital'} <= set(wb.sheetnames)
    ws = wb['Operaciones']
    assert ws['H2'].value == 0.4
    assert ws['K2'].value == 0.4
    assert ws['F2'].value == 100.0
