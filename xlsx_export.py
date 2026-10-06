"""
Exporta resultados de backtest a un archivo .xlsx en memoria.

Módulo puro (no conoce Telegram ni Bitunix): recibe los datos ya
calculados por backtest.py y arma un Workbook de openpyxl. Pensado para
adjuntar el resultado como documento descargable desde el bot de
Telegram, pero reutilizable desde cualquier otro consumidor.
"""

from __future__ import annotations

from datetime import datetime, timezone
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from backtest import BacktestSignal

FONT_NAME = "Arial"
HEADER_FILL = PatternFill(start_color="1F2937", end_color="1F2937", fill_type="solid")
HEADER_FONT = Font(name=FONT_NAME, bold=True, color="FFFFFF")
TITLE_FONT = Font(name=FONT_NAME, bold=True, size=14)
LABEL_FONT = Font(name=FONT_NAME, bold=True)
LONG_FILL = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
SHORT_FILL = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")


def _ms_to_dt(ms: int) -> datetime:
    return datetime.fromtimestamp(ms / 1000, tz=timezone.utc).replace(tzinfo=None)


def _style_header_row(ws, row: int, n_cols: int) -> None:
    for col in range(1, n_cols + 1):
        cell = ws.cell(row=row, column=col)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")


def _autosize(ws, widths: list[int]) -> None:
    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width


def _write_summary_sheet(
    ws,
    symbol: str,
    interval: str,
    horizon: int,
    n_candles: int,
    first_date: datetime,
    last_date: datetime,
    stats: dict,
) -> None:
    ws.title = "Resumen"

    ws["A1"] = f"Backtest {symbol} — {interval}"
    ws["A1"].font = TITLE_FONT
    ws.merge_cells("A1:D1")

    meta = [
        ("Símbolo", symbol),
        ("Intervalo", interval),
        ("Horizonte (velas)", horizon),
        ("Velas analizadas", n_candles),
        ("Desde", first_date.strftime("%Y-%m-%d %H:%M")),
        ("Hasta", last_date.strftime("%Y-%m-%d %H:%M")),
        ("Generado (UTC)", datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M")),
    ]
    row = 3
    for label, value in meta:
        ws.cell(row=row, column=1, value=label).font = LABEL_FONT
        ws.cell(row=row, column=2, value=value)
        row += 1

    row += 1
    table_start = row
    headers = ["Métrica", "Total", "LONG", "SHORT"]
    for col, text in enumerate(headers, start=1):
        ws.cell(row=table_start, column=col, value=text)
    _style_header_row(ws, table_start, len(headers))

    metric_rows = [
        ("Señales totales", "total", "{:d}"),
        ("Con resultado (ya cumplieron el horizonte)", "con_resultado", "{:d}"),
        ("Aciertos", "aciertos", "{:d}"),
        ("Win rate", "win_rate", "pct"),
        ("Cambio medio", "cambio_medio_pct", "pct"),
    ]
    for offset, (label, key, fmt) in enumerate(metric_rows, start=1):
        r = table_start + offset
        ws.cell(row=r, column=1, value=label).font = LABEL_FONT
        for col, group in ((2, "total"), (3, "long"), (4, "short")):
            value = stats[group][key]
            cell = ws.cell(row=r, column=col)
            if value is None:
                cell.value = "—"
            elif fmt == "pct":
                cell.value = round(value, 2)
                cell.number_format = '0.00"%"'
            else:
                cell.value = value

    _autosize(ws, [42, 12, 12, 12])
    ws.freeze_panes = "A1"


def _write_signals_sheet(ws, results: list[BacktestSignal]) -> None:
    ws.title = "Señales"
    headers = [
        "Fecha (UTC)",
        "Lado",
        "Precio cierre",
        "%K",
        "%D",
        "Banda inferior",
        "Banda media",
        "Banda superior",
        f"Precio +N velas",
        "% cambio (a favor)",
        "Resultado",
    ]
    for col, text in enumerate(headers, start=1):
        ws.cell(row=1, column=col, value=text)
    _style_header_row(ws, 1, len(headers))

    for i, r in enumerate(results, start=2):
        if r.pct_change is None:
            resultado = "Pendiente"
        elif r.pct_change > 0:
            resultado = "Acierto"
        else:
            resultado = "Fallo"

        values = [
            _ms_to_dt(r.time),
            r.side,
            r.close,
            r.k,
            r.d,
            r.lower,
            r.middle,
            r.upper,
            r.horizon_close,
            round(r.pct_change, 2) if r.pct_change is not None else None,
            resultado,
        ]
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row=i, column=col, value=value)
            if col == 1:
                cell.number_format = "yyyy-mm-dd hh:mm"
            elif col in (3, 6, 7, 8, 9):
                cell.number_format = "0.0000"
            elif col == 10 and value is not None:
                cell.number_format = '+0.00"%";-0.00"%"'

        fill = LONG_FILL if r.side == "LONG" else SHORT_FILL
        ws.cell(row=i, column=2).fill = fill

    _autosize(ws, [17, 7, 13, 9, 9, 13, 13, 13, 15, 16, 11])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(results) + 1}"


def _write_candles_sheet(ws, klines: list[dict]) -> None:
    ws.title = "Velas"
    headers = ["Fecha (UTC)", "Open", "High", "Low", "Close"]
    for col, text in enumerate(headers, start=1):
        ws.cell(row=1, column=col, value=text)
    _style_header_row(ws, 1, len(headers))

    for i, k in enumerate(klines, start=2):
        values = [
            _ms_to_dt(int(k["time"])),
            float(k.get("open", 0)),
            float(k["high"]),
            float(k["low"]),
            float(k["close"]),
        ]
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row=i, column=col, value=value)
            cell.number_format = "yyyy-mm-dd hh:mm" if col == 1 else "0.0000"

    _autosize(ws, [17, 13, 13, 13, 13])
    ws.freeze_panes = "A2"


def build_backtest_workbook(
    symbol: str,
    interval: str,
    horizon: int,
    klines: list[dict],
    results: list[BacktestSignal],
    stats: dict,
) -> BytesIO:
    """
    Arma el workbook completo del backtest con 3 hojas (Resumen, Señales,
    Velas) y lo devuelve como BytesIO, listo para adjuntar como documento.
    """
    first_date = _ms_to_dt(klines[0]["time"])
    last_date = _ms_to_dt(klines[-1]["time"])

    wb = Workbook()
    summary_ws = wb.active
    _write_summary_sheet(summary_ws, symbol, interval, horizon, len(klines), first_date, last_date, stats)

    signals_ws = wb.create_sheet("Señales")
    _write_signals_sheet(signals_ws, results)

    candles_ws = wb.create_sheet("Velas")
    _write_candles_sheet(candles_ws, klines)

    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if cell.font is None or cell.font.name != FONT_NAME:
                    existing = cell.font
                    cell.font = Font(
                        name=FONT_NAME,
                        bold=existing.bold if existing else False,
                        color=existing.color if existing else None,
                        size=existing.size if existing else 11,
                    )

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer
