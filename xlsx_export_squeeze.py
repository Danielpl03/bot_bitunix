"""Exportación Excel del backtest Squeeze Bollinger + Keltner."""
from datetime import datetime, timezone
from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from risk_management import RiskConfig, calculate_position, scenario_pnl_usdt

DARK=PatternFill('solid',fgColor='1F2937'); WHITE=Font(bold=True,color='FFFFFF'); BOLD=Font(bold=True); TITLE=Font(bold=True,size=14)
def build_squeeze_workbook(symbol, interval, horizon, klines, trades, stats, risk_config: RiskConfig):
    wb=Workbook(); ws=wb.active; ws.title='Resumen'
    ws['A1']=f'Squeeze Bollinger + Keltner — {symbol} {interval}'; ws['A1'].font=TITLE; ws.merge_cells('A1:D1')
    first=datetime.fromtimestamp(int(klines[0]['time'])/1000,timezone.utc); last=datetime.fromtimestamp(int(klines[-1]['time'])/1000,timezone.utc)
    meta=[('Símbolo',symbol),('Temporalidad',interval),('Período UTC',f'{first:%Y-%m-%d %H:%M} → {last:%Y-%m-%d %H:%M}'),('Velas cargadas',len(klines)),('Horizonte',horizon if horizon else 'Sin límite'),('Bollinger','20 períodos, 2 desviaciones'),('Keltner','EMA 20, ATR 20 × 1.5'),('Squeeze','Bandas Bollinger dentro de Keltner en vela previa; liberación en vela de señal'),('Confirmación','Cierre fuera de Bollinger y Keltner + momentum de 12 velas a favor'),('Entrada','Apertura de la vela siguiente'),('Stop','Extremo de las últimas 10 velas, incluida señal'),('Objetivo','2R fijo'),('Posiciones simultáneas','Máximo una'),('Riesgo',risk_config.label),('Costes','No incluye comisiones, funding ni deslizamiento')]
    row=3
    for k,v in meta: ws.cell(row,1,k).font=BOLD; ws.cell(row,2,v); row+=1
    row+=1
    for c,h in enumerate(['Métrica','Total','LONG','SHORT'],1): ws.cell(row,c,h).fill=DARK; ws.cell(row,c).font=WHITE
    metrics=[('Operaciones','total'),('Resueltas','resueltas'),('TP','tp'),('SL','sl'),('Horizonte','horizonte'),('Pendientes','pendientes'),('Acierto %','win_rate'),('R netas','r_total'),('R media','r_medio'),('Profit factor','profit_factor'),('Drawdown máximo R','max_drawdown_r'),('Cambio % acumulado','pnl_pct_total')]
    for ri,(label,key) in enumerate(metrics,row+1):
        ws.cell(ri,1,label)
        for col,g in enumerate(['total','long','short'],2):
            val=stats[g][key]; ws.cell(ri,col,'∞' if isinstance(val,float) and val==float('inf') else val)
    for col,w in enumerate([32,24,24,24],1): ws.column_dimensions[get_column_letter(col)].width=w
    op=wb.create_sheet('Operaciones'); headers=['Hora señal UTC','Hora entrada UTC','Lado','Entrada','Stop','Objetivo 2R','BB inferior','BB media','BB superior','KC inferior','KC media EMA','KC superior','Momentum','Riesgo %','Objetivo R','Resultado','Hora salida UTC','Precio salida','Resultado R','Cambio %','Riesgo USDT','Cantidad estimada','Nominal USDT','Margen aislado USDT','Apalancamiento','PnL USDT estimado']
    for c,h in enumerate(headers,1): op.cell(1,c,h).fill=DARK; op.cell(1,c).font=WHITE
    for r,t in enumerate(trades,2):
        pos=calculate_position(t.entry,t.stop_loss,risk_config); pnl=scenario_pnl_usdt(t.r_multiple,pos) if t.r_multiple is not None else None
        vals=[datetime.fromtimestamp(t.time/1000,timezone.utc).replace(tzinfo=None),datetime.fromtimestamp(t.entry_time/1000,timezone.utc).replace(tzinfo=None),t.side,t.entry,t.stop_loss,t.target,t.bb_lower,t.bb_middle,t.bb_upper,t.kc_lower,t.kc_middle,t.kc_upper,t.momentum,t.risk_pct,t.target_r,t.outcome,datetime.fromtimestamp(t.exit_time/1000,timezone.utc).replace(tzinfo=None) if t.exit_time else None,t.exit_price,t.r_multiple,t.pnl_pct,pos.risk_usdt,pos.quantity,pos.notional_usdt,pos.isolated_margin_usdt,pos.leverage,pnl]
        for c,v in enumerate(vals,1): op.cell(r,c,v)
    op.freeze_panes='A2'; op.auto_filter.ref=op.dimensions
    for c in range(1,len(headers)+1): op.column_dimensions[get_column_letter(c)].width=18
    for r in range(2,op.max_row+1):
        op.cell(r,1).number_format='yyyy-mm-dd hh:mm'; op.cell(r,2).number_format='yyyy-mm-dd hh:mm'; op.cell(r,17).number_format='yyyy-mm-dd hh:mm'
    out=BytesIO(); wb.save(out); out.seek(0); return out
