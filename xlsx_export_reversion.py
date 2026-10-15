"""Exportación Excel del backtest Bollinger + Estocástico (reversión)."""
from __future__ import annotations
from datetime import datetime, timezone
from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from risk_management import RiskConfig, calculate_position, scenario_pnl_usdt

DARK = PatternFill("solid", fgColor="1F2937")
LIGHT = PatternFill("solid", fgColor="E5E7EB")
WHITE_BOLD = Font(name="Arial", bold=True, color="FFFFFF")
BOLD = Font(name="Arial", bold=True)
TITLE = Font(name="Arial", bold=True, size=14)


def build_reversion_workbook(symbol, interval, horizon, klines, trades, stats, risk_config: RiskConfig):
    wb = Workbook()
    ws = wb.active
    ws.title = "Resumen"
    ws["A1"] = f"Reversión Bollinger + Estocástico — {symbol} {interval}"
    ws["A1"].font = TITLE
    ws.merge_cells("A1:D1")
    first = datetime.fromtimestamp(int(klines[0]["time"]) / 1000, timezone.utc)
    last = datetime.fromtimestamp(int(klines[-1]["time"]) / 1000, timezone.utc)
    metadata = [
        ("Símbolo", symbol), ("Temporalidad", interval),
        ("Período histórico UTC", f"{first:%Y-%m-%d %H:%M} → {last:%Y-%m-%d %H:%M}"),
        ("Velas cargadas", len(klines)), ("Horizonte máximo", horizon if horizon else "Sin límite"),
        ("Bandas de Bollinger", "Período 20; desviación 2"),
        ("Estocástico lento", "14, 3, 3; sobreventa < 20; sobrecompra > 80"),
        ("Lookback del stop", "10 velas, incluida la vela de señal"),
        ("Posiciones simultáneas", "Máximo una posición abierta por símbolo"),
        ("Entrada", "Apertura de la vela posterior a la señal"),
        ("Stop", "Mínimo/máximo de las últimas 10 velas, incluida la señal"),
        ("Objetivo", "Banda media de Bollinger fijada al cierre de la vela de señal"),
        ("Riesgo teórico", risk_config.label),
        ("Costes", "No incluye comisiones, funding ni deslizamiento"),
    ]
    row = 3
    for key, val in metadata:
        ws.cell(row, 1, key).font = BOLD
        ws.cell(row, 2, val)
        row += 1
    row += 1
    for col, name in enumerate(["Métrica", "Total", "LONG", "SHORT"], 1):
        ws.cell(row, col, name).fill = DARK
        ws.cell(row, col).font = WHITE_BOLD
    metrics = [
        ("Operaciones detectadas", "total"), ("Operaciones resueltas", "resueltas"),
        ("TP", "tp"), ("SL", "sl"), ("Cerradas por horizonte", "horizonte"),
        ("Pendientes", "pendientes"), ("Acierto %", "win_rate"),
        ("R netas acumuladas", "r_total"), ("R media por operación", "r_medio"),
        ("Profit factor (en R)", "profit_factor"), ("Drawdown máximo (R)", "max_drawdown_r"),
        ("Cambio porcentual acumulado (sin costes)", "pnl_pct_total"),
    ]
    for ri, (label, key) in enumerate(metrics, row + 1):
        ws.cell(ri, 1, label)
        for col, group in enumerate(["total", "long", "short"], 2):
            val = stats[group][key]
            if isinstance(val, float) and val == float("inf"):
                val = "∞"
            ws.cell(ri, col, val)
    ws.freeze_panes = "A3"
    for col, width in enumerate([38, 22, 22, 22], 1):
        ws.column_dimensions[get_column_letter(col)].width = width

    op = wb.create_sheet("Operaciones")
    headers = ["Hora señal UTC", "Hora entrada UTC", "Lado", "Entrada", "Stop", "Objetivo banda media",
               "%K", "%D", "Banda inferior señal", "Media señal", "Banda superior señal", "Riesgo %",
               "Objetivo en R", "Resultado", "Hora salida UTC", "Precio salida", "Resultado R",
               "Cambio %", "Riesgo USDT", "Cantidad estimada", "Nominal USDT", "Margen aislado USDT", "Apalancamiento estimado", "PnL USDT estimado"]
    for c, h in enumerate(headers, 1):
        cell = op.cell(1, c, h); cell.fill = DARK; cell.font = WHITE_BOLD
    for r, t in enumerate(trades, 2):
        pos = calculate_position(t.entry, t.stop_loss, risk_config)
        pnl_usdt = scenario_pnl_usdt(t.r_multiple, pos) if t.r_multiple is not None else None
        vals = [datetime.fromtimestamp(t.time/1000, timezone.utc).replace(tzinfo=None),
                datetime.fromtimestamp(t.entry_time/1000, timezone.utc).replace(tzinfo=None), t.side,
                t.entry, t.stop_loss, t.target, t.k, t.d, t.lower, t.middle, t.upper, t.risk_pct,
                t.target_r, t.outcome,
                datetime.fromtimestamp(t.exit_time/1000, timezone.utc).replace(tzinfo=None) if t.exit_time else None,
                t.exit_price, t.r_multiple, t.pnl_pct, pos.risk_usdt, pos.quantity, pos.notional_usdt,
                pos.isolated_margin_usdt, pos.leverage, pnl_usdt]
        for c, val in enumerate(vals, 1):
            op.cell(r, c, val)
    op.freeze_panes = "A2"
    op.auto_filter.ref = op.dimensions
    for i, width in enumerate([20,20,10,15,15,18,10,10,18,18,18,12,12,15,20,15,12,12,14,18,18,20,18,18], 1):
        op.column_dimensions[get_column_letter(i)].width = width
    for r in range(2, op.max_row + 1):
        op.cell(r, 1).number_format = "yyyy-mm-dd hh:mm"
        op.cell(r, 2).number_format = "yyyy-mm-dd hh:mm"
        op.cell(r, 15).number_format = "yyyy-mm-dd hh:mm"

    candles = wb.create_sheet("Velas")
    for c, h in enumerate(["Hora UTC", "Open", "High", "Low", "Close", "Volumen"], 1):
        cell = candles.cell(1, c, h); cell.fill = DARK; cell.font = WHITE_BOLD
    for r, k in enumerate(klines, 2):
        vals = [datetime.fromtimestamp(int(k["time"])/1000, timezone.utc).replace(tzinfo=None),
                float(k["open"]), float(k["high"]), float(k["low"]), float(k["close"]),
                float(k.get("volume", 0) or 0)]
        for c, val in enumerate(vals, 1): candles.cell(r, c, val)
        candles.cell(r, 1).number_format = "yyyy-mm-dd hh:mm"
    candles.freeze_panes = "A2"
    candles.auto_filter.ref = candles.dimensions
    for i, width in enumerate([20,16,16,16,16,18], 1): candles.column_dimensions[get_column_letter(i)].width = width

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output
